import io
import pytest
from datetime import datetime, timedelta
from decimal import Decimal

from openpyxl import load_workbook

from services.sale_service import SaleService
from services.financial_service import FinancialService
from services.product_service import ProductService
from services.user_service import UserService
from services.export_service import ExportService
from database.models import SupplyOrder, SupplyOrderItem, SupplyStatus


async def _prepare_fifo_sales(session):
    await FinancialService.create_financial_settings(
        session,
        cny_to_som_rate=Decimal("10.00"),
        delivery_cost_per_kg=Decimal("0.00"),
    )
    product = await ProductService.create_product(
        session,
        name="Scale",
        base_price_cny=Decimal("5.00"),
        weight_kg=Decimal("0.200"),
    )
    user = await UserService.get_or_create_user(session, telegram_id=2000, username="fifo_seller")

    now = datetime.now()

    order1 = SupplyOrder(
        tracking_number="TRK1",
        order_date=now - timedelta(days=5),
        delivery_date=now - timedelta(days=4),
        status=SupplyStatus.DELIVERED,
    )
    session.add(order1)
    await session.flush()
    session.add(
        SupplyOrderItem(
            supply_order_id=order1.id,
            product_id=product.id,
            quantity=5,
            unit_price_cny=Decimal("5.00"),
            unit_weight_kg=product.weight_kg,
        )
    )

    order2 = SupplyOrder(
        tracking_number="TRK2",
        order_date=now - timedelta(days=3),
        delivery_date=now - timedelta(days=2),
        status=SupplyStatus.DELIVERED,
    )
    session.add(order2)
    await session.flush()
    session.add(
        SupplyOrderItem(
            supply_order_id=order2.id,
            product_id=product.id,
            quantity=5,
            unit_price_cny=Decimal("6.00"),
            unit_weight_kg=product.weight_kg,
        )
    )

    await session.commit()

    sales_batch1 = await SaleService.create_sale(
        session,
        product_id=product.id,
        seller_id=user.id,
        quantity=5,
        sale_price=Decimal("150.00"),
        sale_date=now + timedelta(hours=1),
    )
    assert len(sales_batch1) == 1
    sale1 = sales_batch1[0]

    sales_batch2 = await SaleService.create_sale(
        session,
        product_id=product.id,
        seller_id=user.id,
        quantity=5,
        sale_price=Decimal("160.00"),
        sale_date=now + timedelta(hours=2),
    )
    assert len(sales_batch2) == 1
    sale2 = sales_batch2[0]

    return product, sale1, sale2


@pytest.mark.asyncio
async def test_create_sale_and_profit(transactional_session):
    session = transactional_session

    # Setup dependencies
    await FinancialService.create_financial_settings(session, cny_to_som_rate=Decimal("12.50"), delivery_cost_per_kg=Decimal("350.00"))
    product = await ProductService.create_product(session, name="Blend", base_price_cny=Decimal("20.00"), weight_kg=Decimal("0.500"))
    user = await UserService.get_or_create_user(session, telegram_id=999, username="seller")

    sale_price = Decimal("500.00")
    sales = await SaleService.create_sale(
        session,
        product_id=product.id,
        seller_id=user.id,
        quantity=2,
        sale_price=sale_price,
        sale_date=datetime.now(),
    )
    assert len(sales) == 1
    sale = sales[0]
    assert sale.quantity == 2
    assert sale.sale_price == sale_price
    assert sale.cost_price is not None
    assert sale.profit is not None


@pytest.mark.asyncio
async def test_get_sales_by_range_and_stats(transactional_session):
    session = transactional_session

    await FinancialService.create_financial_settings(session, cny_to_som_rate=Decimal("12.50"), delivery_cost_per_kg=Decimal("350.00"))
    product = await ProductService.create_product(session, name="Toy", base_price_cny=Decimal("5.00"), weight_kg=Decimal("0.100"))
    user = await UserService.get_or_create_user(session, telegram_id=1000, username="seller2")

    await SaleService.create_sale(session, product_id=product.id, seller_id=user.id, quantity=1, sale_price=Decimal("100.00"), sale_date=datetime.now())

    start = datetime.now() - timedelta(days=1)
    end = datetime.now() + timedelta(days=1)

    sales = await SaleService.get_sales_by_date_range(session, start, end)
    assert len(sales) >= 1

    stats = await SaleService.get_seller_statistics(session, start, end)
    assert isinstance(stats, list)
    assert len(stats) >= 1
    assert "seller_name" in stats[0]


@pytest.mark.asyncio
async def test_sale_cost_fifo_batches(transactional_session):
    product, sale1, sale2 = await _prepare_fifo_sales(transactional_session)

    assert sale1.product_id == product.id
    assert sale2.product_id == product.id

    assert sale1.cost_price == Decimal("50.00")
    assert sale2.cost_price == Decimal("60.00")

    assert sale1.profit == Decimal("500.00")
    assert sale2.profit == Decimal("500.00")


@pytest.mark.asyncio
async def test_sale_split_creates_multiple_rows(transactional_session):
    session = transactional_session

    await FinancialService.create_financial_settings(
        session,
        cny_to_som_rate=Decimal("10.00"),
        delivery_cost_per_kg=Decimal("0.00"),
    )
    product = await ProductService.create_product(
        session,
        name="Scale Split",
        base_price_cny=Decimal("5.00"),
        weight_kg=Decimal("0.200"),
    )
    user = await UserService.get_or_create_user(session, telegram_id=3000, username="split_seller")

    now = datetime.now()

    order1 = SupplyOrder(
        tracking_number="TRK-SPLIT-1",
        order_date=now - timedelta(days=3),
        delivery_date=now - timedelta(days=2),
        status=SupplyStatus.DELIVERED,
    )
    session.add(order1)
    await session.flush()
    session.add(
        SupplyOrderItem(
            supply_order_id=order1.id,
            product_id=product.id,
            quantity=3,
            unit_price_cny=Decimal("5.00"),
            unit_weight_kg=product.weight_kg,
        )
    )

    order2 = SupplyOrder(
        tracking_number="TRK-SPLIT-2",
        order_date=now - timedelta(days=1),
        delivery_date=now,
        status=SupplyStatus.DELIVERED,
    )
    session.add(order2)
    await session.flush()
    session.add(
        SupplyOrderItem(
            supply_order_id=order2.id,
            product_id=product.id,
            quantity=2,
            unit_price_cny=Decimal("6.00"),
            unit_weight_kg=product.weight_kg,
        )
    )

    await session.commit()

    sales = await SaleService.create_sale(
        session,
        product_id=product.id,
        seller_id=user.id,
        quantity=5,
        sale_price=Decimal("140.00"),
        sale_date=now + timedelta(hours=1),
    )

    assert len(sales) == 2
    assert sales[0].quantity == 3
    assert sales[0].cost_price == Decimal("50.00")
    assert sales[1].quantity == 2
    assert sales[1].cost_price == Decimal("60.00")


@pytest.mark.asyncio
async def test_export_accounting_uses_fifo_costs(transactional_session):
    _, sale1, sale2 = await _prepare_fifo_sales(transactional_session)

    excel_bytes = await ExportService.export_accounting_excel(transactional_session)

    workbook = load_workbook(io.BytesIO(excel_bytes))
    ws_accounting = workbook["Accounting"]

    rows = list(ws_accounting.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2

    first_cost, second_cost = rows[0][5], rows[1][5]

    assert first_cost == pytest.approx(float(sale1.cost_price))
    assert second_cost == pytest.approx(float(sale2.cost_price))

    assert rows[0][7] == pytest.approx(first_cost * rows[0][3])
    assert rows[1][7] == pytest.approx(second_cost * rows[1][3])


@pytest.mark.asyncio
async def test_sale_split_uses_fallback_for_missing_stock(transactional_session):
    session = transactional_session

    await FinancialService.create_financial_settings(
        session,
        cny_to_som_rate=Decimal("10.00"),
        delivery_cost_per_kg=Decimal("0.00"),
    )
    product = await ProductService.create_product(
        session,
        name="Scale Fallback",
        base_price_cny=Decimal("6.00"),
        weight_kg=Decimal("0.000"),
    )
    user = await UserService.get_or_create_user(session, telegram_id=3500, username="fallback_seller")

    now = datetime.now()

    order = SupplyOrder(
        tracking_number="TRK-FALLBACK-1",
        order_date=now - timedelta(days=2),
        delivery_date=now - timedelta(days=1),
        status=SupplyStatus.DELIVERED,
    )
    session.add(order)
    await session.flush()
    session.add(
        SupplyOrderItem(
            supply_order_id=order.id,
            product_id=product.id,
            quantity=3,
            unit_price_cny=Decimal("4.00"),
            unit_weight_kg=Decimal("0.000"),
        )
    )
    await session.commit()

    sales = await SaleService.create_sale(
        session,
        product_id=product.id,
        seller_id=user.id,
        quantity=5,
        sale_price=Decimal("120.00"),
        sale_date=now,
    )

    assert len(sales) == 2
    quantities = [s.quantity for s in sales]
    costs = [s.cost_price for s in sales]

    assert sorted(quantities) == [2, 3]
    assert Decimal("40.00") in costs  # from supply batch (4 CNY * 10)
    assert Decimal("60.00") in costs  # fallback (base 6 CNY * 10)


@pytest.mark.asyncio
async def test_export_accounting_with_split_sales(transactional_session):
    session = transactional_session

    await FinancialService.create_financial_settings(
        session,
        cny_to_som_rate=Decimal("10.00"),
        delivery_cost_per_kg=Decimal("0.00"),
    )
    product = await ProductService.create_product(
        session,
        name="Scale Export Split",
        base_price_cny=Decimal("5.00"),
        weight_kg=Decimal("0.200"),
    )
    user = await UserService.get_or_create_user(session, telegram_id=4000, username="export_split_seller")

    now = datetime.now()

    order1 = SupplyOrder(
        tracking_number="TRK-EXP-1",
        order_date=now - timedelta(days=4),
        delivery_date=now - timedelta(days=3),
        status=SupplyStatus.DELIVERED,
    )
    session.add(order1)
    await session.flush()
    session.add(
        SupplyOrderItem(
            supply_order_id=order1.id,
            product_id=product.id,
            quantity=3,
            unit_price_cny=Decimal("5.00"),
            unit_weight_kg=product.weight_kg,
        )
    )

    order2 = SupplyOrder(
        tracking_number="TRK-EXP-2",
        order_date=now - timedelta(days=2),
        delivery_date=now - timedelta(days=1),
        status=SupplyStatus.DELIVERED,
    )
    session.add(order2)
    await session.flush()
    session.add(
        SupplyOrderItem(
            supply_order_id=order2.id,
            product_id=product.id,
            quantity=2,
            unit_price_cny=Decimal("6.00"),
            unit_weight_kg=product.weight_kg,
        )
    )

    await session.commit()

    sales = await SaleService.create_sale(
        session,
        product_id=product.id,
        seller_id=user.id,
        quantity=5,
        sale_price=Decimal("140.00"),
        sale_date=now,
    )
    assert len(sales) == 2

    excel_bytes = await ExportService.export_accounting_excel(session)
    workbook = load_workbook(io.BytesIO(excel_bytes))
    ws_accounting = workbook["Accounting"]

    rows = list(ws_accounting.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2

    quantities = sorted(row[3] for row in rows)
    costs = sorted(row[5] for row in rows)

    assert quantities == [2, 3]
    assert costs == [50.0, 60.0]


@pytest.mark.asyncio
async def test_export_supplies_total_cost_column(transactional_session):
    session = transactional_session
    await _prepare_fifo_sales(session)

    excel_bytes = await ExportService.export_accounting_excel(session)
    workbook = load_workbook(io.BytesIO(excel_bytes))
    ws_supplies = workbook["Поступления"]

    rows = list(ws_supplies.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2

    for row in rows:
        quantity = row[3]
        unit_cost = row[10]
        total_cost = row[11]
        assert total_cost == pytest.approx(unit_cost * quantity)


