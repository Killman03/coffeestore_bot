from datetime import datetime
from decimal import Decimal
from typing import Optional

from sqlalchemy import select
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from database.models import Sale, SupplyOrderItem, SupplyStatus
from services.financial_service import FinancialService


class ExportService:
    """Service for exporting accounting data to Excel."""

    @staticmethod
    async def export_accounting_excel(session: AsyncSession, start: Optional[datetime] = None, end: Optional[datetime] = None) -> bytes:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        wb = Workbook()

        # Sheet 1: Supplies (arrivals)
        ws_supplies = wb.active
        ws_supplies.title = "Поступления"

        supplies_headers = [
            "Дата прибытия",
            "Трек-номер",
            "Товар",
            "Кол-во",
            "Цена закупа",
            "Курс CNY",
            "Цена без доставки (сом)",
            "Кг",
            "Цена за кг",
            "Цена продажи по умолчанию",
            "Себестоимость",
            "Себестоимость поступления",
        ]
        ws_supplies.append(supplies_headers)

        # Query delivered supply items with related order and product
        si_query = (
            select(SupplyOrderItem)
            .options(
                selectinload(SupplyOrderItem.supply_order),
                selectinload(SupplyOrderItem.product),
            )
            .where(SupplyOrderItem.supply_order.has(status=SupplyStatus.DELIVERED))
        )
        si_result = await session.execute(si_query)
        supply_items = si_result.scalars().all()

        # Get default settings in case we can't find settings for a date
        default_settings = await FinancialService.get_or_create_default_settings(session)

        for item in supply_items:
            order = item.supply_order
            product = item.product

            arrival_dt = order.delivery_date or order.order_date
            
            # Get financial settings that were active on the arrival date
            financial_settings = await FinancialService.get_settings_for_date(session, arrival_dt)
            if not financial_settings:
                # Fallback to default settings if no settings found for that date
                financial_settings = default_settings
            
            cny_rate = financial_settings.cny_to_som_rate
            delivery_per_kg = financial_settings.delivery_cost_per_kg
            
            unit_price_cny = item.unit_price_cny if item.unit_price_cny is not None else product.base_price_cny
            unit_weight_kg = item.unit_weight_kg if item.unit_weight_kg is not None else product.weight_kg
            default_sale_price = product.default_sale_price or Decimal(0)

            base_price_som_wo_delivery = unit_price_cny * cny_rate
            cost_price_som = base_price_som_wo_delivery + (unit_weight_kg * delivery_per_kg)

            ws_supplies.append([
                arrival_dt.strftime("%Y-%m-%d"),
                order.tracking_number,
                product.name,
                int(item.quantity),
                float(unit_price_cny),
                float(cny_rate),
                float(base_price_som_wo_delivery),
                float(unit_weight_kg),
                float(delivery_per_kg),
                float(default_sale_price),
                float(cost_price_som),
                float(cost_price_som * Decimal(item.quantity)),
            ])

        # Sheet 2: Sales accounting
        ws_sales = wb.create_sheet(title="Accounting")

        sales_headers = [
            "Дата продажи",
            "Товар",
            "Продавец",
            "Кол-во",
            "Цена продажная (сом)",
            "Себестоимость за ед. (сом)",
            "Цена доставки",
            "Выручка (сом)",
            "Себестоимость (сом)",
            "Прибыль (сом)",
            "Себестоимость Андрей",
            "Себестоимость Женя",
            "Прибыль Андрей",
            "Прибыль Женя",
        ]
        ws_sales.append(sales_headers)

        # Query sales with product and seller eagerly loaded to avoid async lazy loads
        query = select(Sale).options(
            selectinload(Sale.product),
            selectinload(Sale.seller),
        )
        if start:
            query = query.where(Sale.sale_date >= start)
        if end:
            query = query.where(Sale.sale_date <= end)
        result = await session.execute(query.order_by(Sale.sale_date.asc()))
        sales = result.scalars().all()

        settings_cache: dict[tuple[int, int, int], object] = {}

        async def get_settings_for_date_cached(dt: datetime):
            cache_key = (dt.year, dt.month, dt.day)
            if cache_key not in settings_cache:
                settings = await FinancialService.get_settings_for_date(session, dt)
                if not settings:
                    settings = default_settings
                settings_cache[cache_key] = settings
            return settings_cache[cache_key]

        for sale in sales:
            revenue = sale.sale_price * sale.quantity
            total_cost = sale.cost_price * sale.quantity
            profit = sale.profit
            cost_andrey = (total_cost or Decimal(0)) / 2
            cost_jenya = (total_cost or Decimal(0)) / 2
            profit_andrey = (profit or Decimal(0)) / 2
            profit_jenya = (profit or Decimal(0)) / 2

            settings_for_sale = await get_settings_for_date_cached(sale.sale_date)
            delivery_per_kg = settings_for_sale.delivery_cost_per_kg
            product_weight = sale.product.weight_kg if sale.product and sale.product.weight_kg is not None else Decimal(0)
            delivery_price = product_weight * delivery_per_kg

            ws_sales.append([
                sale.sale_date.strftime("%Y-%m-%d %H:%M"),
                sale.product.name,
                sale.seller.full_name,
                int(sale.quantity),
                float(sale.sale_price),
                float(sale.cost_price),
                float(delivery_price),
                float(revenue),
                float(total_cost),
                float(profit),
                float(cost_andrey),
                float(cost_jenya),
                float(profit_andrey),
                float(profit_jenya),
            ])

        # Autosize columns for both sheets
        def autosize(ws):
            from openpyxl.utils import get_column_letter as _gcl
            for col in ws.columns:
                max_length = 0
                column = col[0].column
                for cell in col:
                    try:
                        max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                ws.column_dimensions[_gcl(column)].width = min(max(12, max_length + 2), 50)

        autosize(ws_supplies)
        autosize(ws_sales)

        # Save to bytes
        import io
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return bio.read()


